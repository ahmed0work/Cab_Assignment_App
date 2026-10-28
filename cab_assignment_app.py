# -*- coding: utf-8 -*-
"""
WE Telecom — Cabinet Assignment Tool  (LGX Splitter Edition)
تصميم وتطوير: م/ أحمد محمد السيد — قطاع تخطيط الشبكات غرب القاهرة

هذه النسخة تعمل بنظام LGX Splitter بدلاً من نظام الكاسيت/PLC القديم.
- إعداد على مستوى المشروع: Splitting Ratio  (1:64  أو  1:32)
- نوع الشبكة: معلق (Overhead)  أو  ABF (حفر)
- لا يوجد كاسيت بعد الآن.
"""

import os
import math
import re
import threading

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
    _HAS_TK = True
except Exception:
    tk = None
    _HAS_TK = False

# ── Constants ──────────────────────────────────────────────────────────────
TYPE_ORDER = ['1:4', '1:8', '1:16', '1:32', '1:64']

# خريطة نوع البوكس -> (نوع الـ LGX splitter, سعته)  حسب الـ splitting ratio
# القاعدة: سعة البوكس × سعة السبليتر = عدد العملاء (64 أو 32)
SP_MAP_64 = {
    '1:4':  ('SP_1:16', 16),   # 4  × 16 = 64
    '1:8':  ('SP_1:8',   8),   # 8  ×  8 = 64
    '1:16': ('SP_1:4',   4),   # 16 ×  4 = 64
    '1:32': ('SP_1:2',   2),   # 32 ×  2 = 64
    '1:64': None,              # مباشر — بدون splitter
}
SP_MAP_32 = {
    '1:4':  ('SP_1:8',   8),   # 4  ×  8 = 32
    '1:8':  ('SP_1:4',   4),   # 8  ×  4 = 32
    '1:16': ('SP_1:2',   2),   # 16 ×  2 = 32
    '1:32': None,              # مباشر — بورت واحد
    '1:64': None,              # مباشر — بورتين
}

# عدد الفرعات التي يحتاجها البوكس الواحد في الـ ODF حسب نوع الشبكة
BRANCHES_PER_BOX = {'معلق': 2, 'ABF': 4}

# أقصى عدد بوكسيات على الكابينة (2 × ODF144 = 288 فرعة)
#   معلق: 288 / 2 = 144 بوكس    |    حفر: 288 / 4 = 72 بوكس
MAX_BOXES = {'معلق': 144, 'ABF': 72}

# سعة الفرعات لكل نوع ODF
ODF36_BRANCHES  = 36
ODF144_BRANCHES = 144

SPLITTING_RATIOS = ['1:64', '1:32']

# ── WE Brand Colors ───────────────────────────────────────────────────────
HDR1_COLOR  = '3B0F6E'
HDR2_COLOR  = '6B2FA0'
HDR3_COLOR  = '9B6AC0'
TOTAL_COLOR = 'E8D5F5'
SPARE_COLOR = 'FFF3CD'

# ── GUI WE Palette ─────────────────────────────────────────────────────────
WE_DARK     = '#3B0F6E'
WE_MED      = '#6B2FA0'
WE_LIGHT    = '#9B6AC0'
WE_BG       = '#F3EDF9'
WE_WHITE    = '#FFFFFF'


# ── Core helpers ────────────────────────────────────────────────────────────
def get_sp_map(splitting_ratio):
    """يرجّع خريطة البوكس->السبليتر حسب الـ ratio."""
    return SP_MAP_64 if splitting_ratio == '1:64' else SP_MAP_32


def canonical_col(name):
    """
    يحوّل اسم أي عمود إلى الاسم القياسي بناءً على الكلمات المفتاحية،
    عشان الأعمدة تشتغل بالاسم مش بالترتيب.
    - أي عمود فيه bundle أو duct يعتبر عمود البندل.
    - التمييز بين رقم البندل و نوع/طول البندل بوجود len/type.
    """
    n = str(name).strip().upper().replace(' ', '').replace('-', '_')

    if 'TYPE' in n and 'BOX' in n:
        return 'TYPE_BOX'
    if ('BUNDLE' in n or 'DUCT' in n) and ('LEN' in n or 'TYPE' in n):
        return 'BUNDLETYPE_LENGTH'
    if 'BUNDLE' in n or 'DUCT' in n:
        return 'BUNDLE_NO'
    if 'CAB' in n and ('LEN' in n or 'FROM' in n):
        return 'LENTH_FROM_CAB'
    if 'CAB' in n:
        return 'CAB_NO'
    if 'BOX' in n:
        return 'BOX_NO'
    return None


def parse_way(val):
    """
    يستخرج سعة البندل (7/12/24/48...) من نص النوع مهما كانت اللاحقة:
    '24PPD_230m' -> 24 ، '12W_170m' -> 12 ، '48PPD' -> 48 ، '7W' -> 7
    القاعدة: الرقم الموجود في أول جزء النوع (قبل الشرطة السفلية).
    """
    s = str(val).split('_')[0].strip()
    m = re.match(r'(\d+)', s)
    return int(m.group(1)) if m else 0


def bundle_type_label(val):
    """
    يرجّع نوع البندل/الـ duct بلاحقته الصحيحة زي ما هي في الداتا:
        معلق: '24PPD_230m' -> '24PPD'   (Preloaded Duct)
        حفر : '24Way_230m' -> '24Way'
    مع توحيد شكل اللاحقة (PPD / Way / W) عشان الأنواع المتشابهة تتجمّع.
    يرجّع '' لو مفيش نوع.
    """
    s = str(val).split('_')[0].strip()
    m = re.match(r'(\d+)\s*([A-Za-z]+)', s)
    if not m:
        return ''
    way = int(m.group(1))
    suffix = m.group(2)
    up = suffix.upper()
    if up == 'PPD':
        suffix = 'PPD'
    elif up.startswith('WAY'):
        suffix = 'Way'
    elif up == 'W':
        suffix = 'W'
    else:
        suffix = up
    return f'{way}{suffix}'


def _odf_plan(network_type, total_boxes):
    """
    يحدّد عدد الـ output ODFs (بدون الـ input ODF36 الثابت) بأقل عدد ممكن،
    مبني على السعات الفعلية للفرعات — نفس المنطق للمعلق والحفر:
        - سعة output ODF36  = 36 / branches_per_box   (18 معلق / 9 حفر)
        - سعة output ODF144 = 144 / branches_per_box  (72 معلق / 36 حفر)
    القاعدة (حسب أقل عدد ODFs يكفي):
        B ≤ cap36               -> 1 output ODF36
        B ≤ cap144              -> 1 ODF144
        B ≤ cap36 + cap144      -> 1 output ODF36 + 1 ODF144
        B ≤ 2·cap144            -> 2 ODF144
    يرجّع: (عدد output ODF36 , عدد ODF144).
    """
    bpb    = BRANCHES_PER_BOX[network_type]
    cap36  = ODF36_BRANCHES  // bpb
    cap144 = ODF144_BRANCHES // bpb
    B = total_boxes

    if B <= cap36:
        return 1, 0
    if B <= cap144:
        return 0, 1
    if B <= cap36 + cap144:
        return 1, 1
    if B <= 2 * cap144:
        return 0, 2
    # تخطّى الحد الأقصى — نوفّر ODF144 كافية (حماية فقط)
    return 0, math.ceil(B / cap144)


def odf_counts(network_type, total_boxes):
    """
    إجمالي عدد الـ ODF36 و ODF144 (شامل الـ input ODF36 الثابت).
    """
    out36, out144 = _odf_plan(network_type, total_boxes)
    return out36 + 1, out144      # +1 = الـ input ODF36 الثابت


def output_odf_list(network_type, total_boxes):
    """
    قائمة الـ output ODFs اللي هتتوصل عليها البوكسيات (بدون الـ input ODF36).
    ترتيب الملء: ODF36 (لو فيه output) ثم ODF144.
    كل عنصر: (اسم الـ ODF, سعته بالبوكسيات).
    يوجد ضمان بأن السعة الكلية تكفي كل البوكسيات (حماية للحالات الحدية).
    """
    out36, out144 = _odf_plan(network_type, total_boxes)
    bpb    = BRANCHES_PER_BOX[network_type]
    cap36  = ODF36_BRANCHES  // bpb     # سعة الـ output ODF36 بالبوكسيات
    cap144 = ODF144_BRANCHES // bpb     # سعة الـ output ODF144 بالبوكسيات

    plan = []
    plan += [('ODF36', cap36)] * out36
    plan += [('ODF144', cap144)] * out144

    # حماية: لو السعة مش كافية (حالة حدية في الحفر) نزوّد ODF144
    total_cap = sum(c for _, c in plan)
    while total_cap < total_boxes:
        plan.append(('ODF144', cap144))
        total_cap += cap144
    if not plan:                        # fallback نظري
        plan.append(('ODF144', cap144))

    # تسمية النسخ: ODF36-1 / ODF144-1 / ODF144-2 ...
    named, c36, c144 = [], 0, 0
    for kind, cap in plan:
        if kind == 'ODF36':
            c36 += 1
            named.append((f'ODF36-{c36}', cap))
        else:
            c144 += 1
            named.append((f'ODF144-{c144}', cap))
    return named


def build_cab_rows(cab_df, network_type, splitting_ratio):
    """
    يبني صفوف تاب الكابينة:
    - يرتّب البوكسيات حسب النوع (1:4 -> 1:8 -> ...) ثم BOX_NO.
    - يجمّع بوكسيات نفس النوع على السبليتر المناسب لحد ما يتملّى، ثم سبليتر جديد.
    - يوصّل كل بوكس على فرعات ODF (2 معلق / 4 حفر) ويملأ الـ output ODFs بالترتيب.
    - يولّد صفوف spare (صفراء) للبورتات الفاضية في أي سبليتر مش مكتمل.
    """
    sp_map = get_sp_map(splitting_ratio)
    bpb    = BRANCHES_PER_BOX[network_type]

    total_boxes = len(cab_df)
    odf_list = output_odf_list(network_type, total_boxes)

    df = cab_df.copy()
    df['type_rank'] = df['TYPE_BOX'].apply(
        lambda t: TYPE_ORDER.index(t) if t in TYPE_ORDER else 99)
    df = df.sort_values(['type_rank', 'BOX_NO']).reset_index(drop=True)

    rows = []

    # حالة السبليتر (متسلسل على مستوى الكابينة)
    sp_num      = 0
    cur_sp_key  = None    # نوع البوكس الخاص بالسبليتر الحالي
    cur_sp_lbl  = None
    cur_sp_cap  = None
    cur_sp_port = 0

    # حالة الـ ODF
    odf_idx        = 0
    odf_branch_used = 0

    def emit_spares():
        """صفوف صفراء للبورتات الفاضية في السبليتر الحالي."""
        if cur_sp_lbl is not None and cur_sp_cap is not None:
            for p in range(cur_sp_port + 1, cur_sp_cap + 1):
                rows.append({
                    'BOX_NO': '', 'BOX_TYPE': '', 'BUNDLE_NO': '',
                    'ODF': '', 'ODF_BRANCH': '',
                    'SP_TYPE': cur_sp_lbl, 'SP_NO_PORT': f'{sp_num}/{p:02d}',
                    'is_spare': True,
                })

    for _, r in df.iterrows():
        btype   = r['TYPE_BOX']
        mapping = sp_map.get(btype, None)

        # ── تخصيص السبليتر ──
        if mapping is None:
            # بوكس مباشر: نقفل أي سبليتر مفتوح (بفراغاته) ونفضّي الحالة
            emit_spares()
            cur_sp_key = cur_sp_lbl = cur_sp_cap = None
            cur_sp_port = 0
            sp_lbl_out, sp_port_out = '', ''
        else:
            lbl, cap = mapping
            if btype != cur_sp_key:
                emit_spares()               # اقفل السبليتر السابق (لو مش مكتمل)
                sp_num += 1
                cur_sp_key, cur_sp_lbl, cur_sp_cap = btype, lbl, cap
                cur_sp_port = 1
            else:
                if cur_sp_port >= cap:       # السبليتر اتملّى -> سبليتر جديد بنفس النوع
                    sp_num += 1
                    cur_sp_port = 1
                else:
                    cur_sp_port += 1
            sp_lbl_out  = lbl
            sp_port_out = f'{sp_num}/{cur_sp_port:02d}'

        # ── تخصيص فرعات الـ ODF (كل بوكس بياخد فرعات) ──
        name, cap_boxes = odf_list[odf_idx]
        cap_branches = cap_boxes * bpb
        if odf_branch_used + bpb > cap_branches:
            odf_idx += 1
            odf_branch_used = 0
            name, cap_boxes = odf_list[odf_idx]
            cap_branches = cap_boxes * bpb
        start_branch = odf_branch_used + 1
        end_branch   = odf_branch_used + bpb
        odf_branch_used = end_branch

        rows.append({
            'BOX_NO':     r['BOX_NO'],
            'BOX_TYPE':   btype,
            'BUNDLE_NO':  r['BUNDLE_NO'],
            'ODF':        name,
            'ODF_BRANCH': f'ODF_{start_branch}:{end_branch}',
            'SP_TYPE':    sp_lbl_out,
            'SP_NO_PORT': sp_port_out,
            'is_spare':   False,
        })

    emit_spares()   # اقفل آخر سبليتر
    return rows


def compute_cab_stats(cab_df, rows, network_type, splitting_ratio, unique_bundle_types=None):
    """يحسب إحصائيات الكابينة (بدون صفوف الـ spare)."""
    import pandas as pd
    if unique_bundle_types is None:
        unique_bundle_types = []

    actual = [r for r in rows if not r.get('is_spare', False)]
    total_boxes   = len(actual)
    boxes_by_type = {t: sum(1 for r in actual if r['BOX_TYPE'] == t) for t in TYPE_ORDER}
    bundles_count = cab_df['BUNDLE_NO'].nunique()

    # ── السبليترز (مشتقّة من صفوف التاب لضمان التطابق) ──
    sp_groups = {}   # sp_num -> {'label':.., 'max_port':..}
    for r in actual:
        if r.get('SP_NO_PORT') and r.get('SP_TYPE') and '/' in r['SP_NO_PORT']:
            num, port = r['SP_NO_PORT'].split('/')
            g = sp_groups.setdefault(num, {'label': r['SP_TYPE'], 'max_port': 0})
            g['max_port'] = max(g['max_port'], int(port))

    total_splitters = len(sp_groups)
    sp_by_type = {'SP_1:16': 0, 'SP_1:8': 0, 'SP_1:4': 0, 'SP_1:2': 0}
    for g in sp_groups.values():
        if g['label'] in sp_by_type:
            sp_by_type[g['label']] += 1

    total_welding = sum(g['max_port'] + 1 for g in sp_groups.values())

    # ── الـ ODF ──
    odfs_used = len(set(r['ODF'] for r in actual if r.get('ODF')))
    odf36, odf144 = odf_counts(network_type, total_boxes)

    # ── Patch Cords / Rack / Ports ──
    patch_cords = total_boxes + total_splitters
    n_rack = 2 if total_splitters > 12 else 1

    if splitting_ratio == '1:64':
        used_ports = total_splitters + boxes_by_type.get('1:64', 0)
    else:  # 1:32
        used_ports = (total_splitters
                      + boxes_by_type.get('1:32', 0)
                      + 2 * boxes_by_type.get('1:64', 0))

    # ── الأطوال القديمة (محتفظ بيها) ──
    # 4F Length: تُحسب في الحفر (ABF) فقط — في المعلق تبقى صفر بالتصميم
    four_f_length = 0
    if network_type == 'ABF' and 'LENTH_FROM_CAB' in cab_df.columns:
        four_f_length = pd.to_numeric(cab_df['LENTH_FROM_CAB'], errors='coerce').fillna(0).sum()
    singles_length = total_boxes * 15

    bundle_lengths = {t: 0.0 for t in unique_bundle_types}
    if 'BUNDLETYPE_LENGTH' in cab_df.columns:
        uniq = cab_df.dropna(subset=['BUNDLE_NO']).drop_duplicates(subset=['BUNDLE_NO'])
        for _, row in uniq.iterrows():
            b_str = str(row.get('BUNDLETYPE_LENGTH', '')).strip()
            b_type = bundle_type_label(b_str)
            parts = b_str.split('_')
            b_len = 0.0
            if len(parts) >= 2:
                try:
                    b_len = float(parts[1].lower().replace('m', '').strip())
                except ValueError:
                    b_len = 0.0
            if b_type in bundle_lengths:
                bundle_lengths[b_type] += b_len

    # ── End Cab (في الحفر ABF فقط) ──
    end_cab = 0
    if network_type == 'ABF' and 'BUNDLETYPE_LENGTH' in cab_df.columns:
        total_capacity = 0
        for _bno, grp in cab_df.groupby('BUNDLE_NO'):
            way = 0
            for v in grp['BUNDLETYPE_LENGTH'].dropna():
                w = parse_way(v)
                if w:
                    way = w
                    break
            total_capacity += way
        end_cab = total_capacity - total_boxes

    return {
        'total_boxes':     total_boxes,
        'bundles_count':   bundles_count,
        'odfs_used':       odfs_used,
        'boxes_by_type':   boxes_by_type,
        'total_splitters': total_splitters,
        'sp_by_type':      sp_by_type,
        'patch_cords':     patch_cords,
        'n_rack':          n_rack,
        'odf36':           odf36,
        'odf144':          odf144,
        'used_ports':      used_ports,
        'total_welding':   total_welding,
        '4f_length':       four_f_length,
        'singles_length':  singles_length,
        'end_cab':         end_cab,
        'bundle_lengths':  bundle_lengths,
    }


# ── Excel styles ────────────────────────────────────────────────────────────
def _make_excel_styles():
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    H1  = PatternFill('solid', start_color=HDR1_COLOR)
    H2  = PatternFill('solid', start_color=HDR2_COLOR)
    H3  = PatternFill('solid', start_color=HDR3_COLOR)
    TOT = PatternFill('solid', start_color=TOTAL_COLOR)
    WF  = Font(bold=True, color='FFFFFF')
    TW  = Side(style='thin', color='FFFFFF')
    BW  = Border(left=TW, right=TW, top=TW, bottom=TW)
    CTR = Alignment(horizontal='center', vertical='center')
    DB  = Border(left=Side(style='thin'), right=Side(style='thin'),
                 top=Side(style='thin'),  bottom=Side(style='thin'))
    return H1, H2, H3, TOT, WF, TW, BW, CTR, DB


def write_cab_sheet(wb, cab_name, rows, network_type, splitting_ratio):
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import PatternFill as _PF

    sheet_title = str(cab_name)[:31]
    if sheet_title in wb.sheetnames:
        del wb[sheet_title]
    ws = wb.create_sheet(title=sheet_title)

    H1, H2, H3, TOT, WF, TW, BW, CTR, DB = _make_excel_styles()

    def hcell(row, col, value, fill, merge_r=None, merge_c=None):
        c = ws.cell(row=row, column=col, value=value)
        c.fill = fill; c.font = WF; c.alignment = CTR; c.border = BW
        if merge_r or merge_c:
            ws.merge_cells(start_row=row, start_column=col,
                           end_row=merge_r or row, end_column=merge_c or col)

    hcell(1, 1,
          f'Cab Name:   {cab_name}      |      Network: {network_type}      |      '
          f'Splitting Ratio: {splitting_ratio}',
          H1, merge_c=7)

    hcell(2, 1, 'Box Data',            H2, merge_c=3)
    hcell(2, 4, 'ODF Assignment',      H2, merge_c=5)
    hcell(2, 6, 'Splitter Assignment', H2, merge_c=7)

    for col, val in [(1, 'Box Number'), (2, 'Box Type'), (3, 'Bundle Number'),
                     (4, 'ODF'), (5, 'ODF Branches'),
                     (6, 'Splitter Type'), (7, 'SP_NO / SP_PORT')]:
        hcell(3, col, val, H3)

    SPARE_FILL = _PF('solid', start_color=SPARE_COLOR)
    for r_idx, r in enumerate(rows, 4):
        is_spare = r.get('is_spare', False)
        for c_idx, val in enumerate([
            r['BOX_NO'], r['BOX_TYPE'], r['BUNDLE_NO'],
            r['ODF'], r['ODF_BRANCH'], r['SP_TYPE'], r['SP_NO_PORT'],
        ], 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.alignment = CTR
            cell.border    = DB
            if is_spare:
                cell.fill = SPARE_FILL

    for i, w in enumerate([13, 11, 14, 12, 16, 15, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ── Summary column model (shared by the two summary sheets) ─────────────────
def _summary_headers(unique_bundle_types):
    base = [
        'Total Boxes', 'Bundles',
        '1:4 Boxes', '1:8 Boxes', '1:16 Boxes', '1:32 Boxes', '1:64 Boxes',
        'SP_1:2', 'SP_1:4', 'SP_1:8', 'SP_1:16', 'Total num of SP',
        'Patch Cords', '#Rack', 'ODF36', 'ODF144',
        'Used Ports', 'Total Welding', 'End Cab',
        '4F Length', 'Singles Length',
    ]
    for t in unique_bundle_types:
        base.append(f'{t} Total Length (m)')
    return base


def _summary_row_values(s, unique_bundle_types):
    vals = [
        s['total_boxes'], s['bundles_count'],
        s['boxes_by_type'].get('1:4',  0),
        s['boxes_by_type'].get('1:8',  0),
        s['boxes_by_type'].get('1:16', 0),
        s['boxes_by_type'].get('1:32', 0),
        s['boxes_by_type'].get('1:64', 0),
        s['sp_by_type'].get('SP_1:2',  0),
        s['sp_by_type'].get('SP_1:4',  0),
        s['sp_by_type'].get('SP_1:8',  0),
        s['sp_by_type'].get('SP_1:16', 0),
        s['total_splitters'],
        s['patch_cords'], s['n_rack'], s['odf36'], s['odf144'],
        s['used_ports'], s['total_welding'], s['end_cab'],
        s.get('4f_length', 0), s.get('singles_length', 0),
    ]
    for t in unique_bundle_types:
        vals.append(s.get('bundle_lengths', {}).get(t, 0.0))
    return vals


def write_cabinets_summary(wb, cab_stats_list, unique_bundle_types):
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    if 'Cabinets Summary' in wb.sheetnames:
        del wb['Cabinets Summary']
    ws = wb.create_sheet(title='Cabinets Summary')

    H1, H2, H3, TOT, WF, TW, BW, CTR, DB = _make_excel_styles()
    headers = _summary_headers(unique_bundle_types)
    TOTAL_COLS = 1 + len(headers)

    def hcell(row, col, value, fill, merge_r=None, merge_c=None):
        c = ws.cell(row=row, column=col, value=value)
        c.fill = fill; c.font = WF; c.alignment = CTR; c.border = BW
        if merge_r or merge_c:
            ws.merge_cells(start_row=row, start_column=col,
                           end_row=merge_r or row, end_column=merge_c or col)

    hcell(1, 1, 'Cabinets Summary', H1, merge_c=TOTAL_COLS)
    hcell(2, 1,  'Cabinet',       H2, merge_r=3)
    hcell(2, 2,  'General',       H2, merge_c=3)
    hcell(2, 4,  'Boxes by Type', H2, merge_c=8)
    hcell(2, 9,  'LGX Splitters', H2, merge_c=13)
    hcell(2, 14, 'Racks & ODFs',  H2, merge_c=17)
    hcell(2, 18, 'Connections',   H2, merge_c=20)
    hcell(2, 21, 'Lengths',       H2, merge_c=TOTAL_COLS)

    for i, h in enumerate(headers, 2):
        hcell(3, i, h, H3)

    totals = [0] * (TOTAL_COLS + 1)
    for r_idx, (cab_name, s) in enumerate(cab_stats_list, 4):
        vals = [cab_name] + _summary_row_values(s, unique_bundle_types)
        for c_idx, val in enumerate(vals, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.alignment = CTR; cell.border = DB
            if c_idx > 1 and isinstance(val, (int, float)):
                totals[c_idx] += val

    tot_r = len(cab_stats_list) + 4
    tf = Font(bold=True, color='3B0F6E')
    # #Rack و ODF36/144 لا يُجمعوا بشكل يدوي عادةً لكن نجمعهم كإجمالي مشروع
    total_vals = ['TOTAL'] + [totals[i] for i in range(2, TOTAL_COLS + 1)]
    for c_idx, val in enumerate(total_vals, 1):
        cell = ws.cell(row=tot_r, column=c_idx, value=val)
        cell.font = tf; cell.fill = TOT; cell.alignment = CTR; cell.border = DB

    widths = ([16, 11, 9] + [10] * 5 + [9] * 4
              + [13, 13, 8, 9, 9, 11, 13, 10, 12, 14]
              + [18] * len(unique_bundle_types))
    for i, w in enumerate(widths[:TOTAL_COLS], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_project_summary(wb, cab_stats_list, unique_bundle_types):
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    if 'Project Summary' in wb.sheetnames:
        del wb['Project Summary']
    ws = wb.create_sheet(title='Project Summary')

    H1, H2, H3, TOT, WF, TW, BW, CTR, DB = _make_excel_styles()
    headers = _summary_headers(unique_bundle_types)
    TOTAL_COLS = 1 + len(headers)

    def hcell(row, col, value, fill, merge_r=None, merge_c=None):
        c = ws.cell(row=row, column=col, value=value)
        c.fill = fill; c.font = WF; c.alignment = CTR; c.border = BW
        if merge_r or merge_c:
            ws.merge_cells(start_row=row, start_column=col,
                           end_row=merge_r or row, end_column=merge_c or col)

    hcell(1, 1, f'Project Summary  —  {len(cab_stats_list)} Cabinets', H1, merge_c=TOTAL_COLS)
    hcell(2, 1,  'Metric',        H2, merge_r=3)
    hcell(2, 2,  'General',       H2, merge_c=3)
    hcell(2, 4,  'Boxes by Type', H2, merge_c=8)
    hcell(2, 9,  'LGX Splitters', H2, merge_c=13)
    hcell(2, 14, 'Racks & ODFs',  H2, merge_c=17)
    hcell(2, 18, 'Connections',   H2, merge_c=20)
    hcell(2, 21, 'Lengths',       H2, merge_c=TOTAL_COLS)

    for i, h in enumerate(headers, 2):
        hcell(3, i, h, H3)

    # جمع كل الكابائن
    agg = None
    for _, s in cab_stats_list:
        vals = _summary_row_values(s, unique_bundle_types)
        if agg is None:
            agg = list(vals)
        else:
            agg = [a + b for a, b in zip(agg, vals)]
    if agg is None:
        agg = [0] * len(headers)

    row_vals = ['PROJECT TOTAL'] + agg
    tf = Font(bold=True, color='3B0F6E')
    for c_idx, val in enumerate(row_vals, 1):
        cell = ws.cell(row=4, column=c_idx, value=val)
        cell.font = tf; cell.fill = TOT; cell.alignment = CTR; cell.border = DB

    widths = ([16, 11, 9] + [10] * 5 + [9] * 4
              + [13, 13, 8, 9, 9, 11, 13, 10, 12, 14]
              + [18] * len(unique_bundle_types))
    for i, w in enumerate(widths[:TOTAL_COLS], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ── Main processing ─────────────────────────────────────────────────────────
def process_file(input_path, output_path, network_type, splitting_ratio, log_callback):
    import pandas as pd
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter as gcl
    import tempfile

    log_callback("📂 Reading file...")
    log_callback(f"🔧 Network type    : {network_type}")
    log_callback(f"🔀 Splitting Ratio : {splitting_ratio}")
    log_callback(f"🌿 Branches / Box  : {BRANCHES_PER_BOX[network_type]}\n")

    # تحويل .xls -> .xlsx لو لزم
    if input_path.lower().endswith('.xls'):
        log_callback("🔄 Converting .xls to .xlsx format...")
        tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        tmp.close()
        xls_data = pd.read_excel(input_path, sheet_name=None)
        with pd.ExcelWriter(tmp.name, engine='openpyxl') as writer:
            for sheet_name, sheet_df in xls_data.items():
                sheet_df.to_excel(writer, sheet_name=sheet_name, index=False)
        input_path = tmp.name
        cleanup_tmp, tmp_path = True, tmp.name
    else:
        cleanup_tmp, tmp_path = False, None

    if 'Summary' not in load_workbook(input_path, read_only=True).sheetnames:
        raise ValueError('لم يتم العثور على ورقة باسم "Summary" (بحرف S كبير).')

    df = pd.read_excel(input_path, sheet_name='Summary')

    # ── قراءة الأعمدة بالاسم (مش بالترتيب) ──
    rename_map = {}
    for col in df.columns:
        canon = canonical_col(col)
        if canon and canon not in rename_map.values():
            rename_map[col] = canon
    df = df.rename(columns=rename_map)

    required = ['CAB_NO', 'BOX_NO', 'BUNDLE_NO', 'TYPE_BOX']
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError(f"أعمدة ناقصة أو غير معروفة: {', '.join(missing)}")

    # تنظيف الأنواع
    df['TYPE_BOX'] = df['TYPE_BOX'].astype(str).str.strip()
    df = df.dropna(subset=['CAB_NO', 'BOX_NO'])
    try:
        df['BOX_NO'] = pd.to_numeric(df['BOX_NO'], errors='coerce')
    except Exception:
        pass

    # أنواع البنادل الفريدة (للأطوال)
    unique_bundle_types = set()
    if 'BUNDLETYPE_LENGTH' in df.columns:
        for val in df['BUNDLETYPE_LENGTH'].dropna().unique():
            t = bundle_type_label(val)
            if t:
                unique_bundle_types.add(t)
    unique_bundle_types = sorted(
        unique_bundle_types, key=lambda x: parse_way(x), reverse=True)

    # ── ورقة Extracted Data ──
    wb = load_workbook(input_path)
    wb['Summary'].title = 'Extracted Data'
    for sh in [s for s in wb.sheetnames if s != 'Extracted Data']:
        del wb[sh]

    ws_ed = wb['Extracted Data']
    PURP = PatternFill('solid', start_color=HDR2_COLOR)
    WF   = Font(bold=True, color='FFFFFF')
    THIN = Side(style='thin', color='FFFFFF')
    BRD  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    CTR  = Alignment(horizontal='center', vertical='center')
    DATA_BRD = Border(left=Side(style='thin'), right=Side(style='thin'),
                      top=Side(style='thin'),  bottom=Side(style='thin'))
    for cell in ws_ed[1]:
        if cell.value is not None:
            cell.fill = PURP; cell.font = WF; cell.alignment = CTR; cell.border = BRD
            ws_ed.column_dimensions[gcl(cell.column)].width = 16
    for row in ws_ed.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = CTR; cell.border = DATA_BRD

    cabs = sorted(df['CAB_NO'].unique(), key=lambda x: str(x))
    log_callback(f"✅ Found {len(cabs)} cabinets, {len(df)} boxes total\n")

    cab_stats_list = []
    for cab_name in cabs:
        cab_df = df[df['CAB_NO'] == cab_name].copy()
        rows   = build_cab_rows(cab_df, network_type, splitting_ratio)
        write_cab_sheet(wb, cab_name, rows, network_type, splitting_ratio)
        stats  = compute_cab_stats(cab_df, rows, network_type,
                                   splitting_ratio, unique_bundle_types)
        cab_stats_list.append((cab_name, stats))
        _max = MAX_BOXES.get(network_type, 144)
        if stats['total_boxes'] > _max:
            log_callback(
                f"  ⚠️ {cab_name}: {stats['total_boxes']} box يتخطى الحد الأقصى "
                f"({_max} لـ {network_type}) — تمت إضافة ODF144 إضافي لاستيعاب الزيادة.")
        log_callback(
            f"  ✔ {cab_name} → {stats['total_boxes']} box | "
            f"{stats['total_splitters']} SP | {stats['patch_cords']} patch | "
            f"{stats['n_rack']} rack | ODF36={stats['odf36']} ODF144={stats['odf144']}"
            + (f" | EndCab={stats['end_cab']}" if network_type == 'ABF' else ""))

    log_callback("\n📊 Building Cabinets Summary...")
    write_cabinets_summary(wb, cab_stats_list, unique_bundle_types)
    log_callback("📊 Building Project Summary...")
    write_project_summary(wb, cab_stats_list, unique_bundle_types)

    desired_order = (['Extracted Data', 'Cabinets Summary', 'Project Summary'] +
                     [str(c)[:31] for c in cabs])
    existing = set(wb.sheetnames)
    wb._sheets = [wb[s] for s in desired_order if s in existing]

    wb.save(output_path)
    log_callback(f"\n💾 Saved workbook: {output_path}")

    # ── تصدير PDF لكل كابينة (ويندوز فقط) ──
    log_callback("\n🗂️ Creating PDF files for each cabinet...")
    wb_com = None
    try:
        import win32com.client
        import pythoncom
        pythoncom.CoInitialize()

        base_dir  = os.path.dirname(os.path.abspath(output_path))
        base_name = os.path.splitext(os.path.basename(output_path))[0]
        pdf_folder = os.path.normpath(os.path.join(base_dir, f"{base_name}_PDFs"))
        os.makedirs(pdf_folder, exist_ok=True)
        log_callback(f"📁 Target folder: {pdf_folder}")

        excel_app = win32com.client.DispatchEx("Excel.Application")
        excel_app.Interactive = False
        excel_app.Visible = False
        excel_app.DisplayAlerts = False

        abs_output_path = os.path.normpath(os.path.abspath(output_path))
        wb_com = excel_app.Workbooks.Open(abs_output_path)

        for cab_name in cabs:
            try:
                ws_com = wb_com.Worksheets(str(cab_name)[:31])
                pdf_file_path = os.path.normpath(os.path.join(pdf_folder, f"{cab_name}.pdf"))
                ws_com.PageSetup.Zoom = False
                ws_com.PageSetup.FitToPagesWide = 1
                ws_com.PageSetup.FitToPagesTall = False
                ws_com.PageSetup.Orientation = 1
                ws_com.ExportAsFixedFormat(0, pdf_file_path)
                log_callback(f"  📄 Exported PDF: {cab_name}.pdf")
            except Exception as e_inner:
                log_callback(f"  ⚠️ Skipped {cab_name}: {str(e_inner)}")

        wb_com.Close(False)
        excel_app.Quit()
        pythoncom.CoUninitialize()
        log_callback("✅ All PDFs generated successfully!")
    except Exception as e:
        log_callback(f"ℹ️ PDF export skipped ({str(e)}).")
        try:
            if wb_com is not None:
                wb_com.Close(False)
                excel_app.Quit()
        except Exception:
            pass

    if cleanup_tmp and tmp_path:
        try:
            os.unlink(tmp_path)
        except Exception:
            pass
    log_callback(f"\n🎉 Done! Saved to:\n{output_path}")


# ── GUI Palette ────────────────────────────────────────────────────────────
GOLD        = '#FFD700'
CARD_BG     = '#FFFFFF'
CARD_BORDER = '#D8C8EE'
ENTRY_BG    = '#F8F4FC'
DIVIDER     = '#DDD0EE'
LOG_BG      = '#0D0D1A'
LOG_FG      = '#00E676'
LOG_HDR_BG  = '#1A1030'


def _card(parent, **kw):
    outer = tk.Frame(parent, bg=CARD_BORDER, bd=0)
    inner = tk.Frame(outer, bg=CARD_BG, **kw)
    inner.pack(fill='both', expand=True, padx=1, pady=1)
    return outer, inner


_TkBase = tk.Tk if _HAS_TK else object


class App(_TkBase):
    def __init__(self):
        super().__init__()
        self.title("WE Telecom — Cabinet Assignment Tool (LGX)")
        self.configure(bg='#EDE7F6')

        self.update_idletasks()
        sw = self.winfo_screenwidth()
        sh = self.winfo_screenheight()
        WIN_W = min(860, sw - 40)
        WIN_H = min(760, sh - 70)
        x = (sw - WIN_W) // 2
        y = max(10, (sh - WIN_H) // 2 - 20)
        self.geometry(f"{WIN_W}x{WIN_H}+{x}+{y}")
        self.minsize(WIN_W, 560)
        self.maxsize(WIN_W, sh - 60)
        self.resizable(False, True)

        self.input_path      = tk.StringVar()
        self.output_path     = tk.StringVar()
        self.network_type    = tk.StringVar(value='معلق')
        self.splitting_ratio = tk.StringVar(value='1:64')

        self._build_ui()

    def _section_header(self, parent, icon, title):
        bar = tk.Frame(parent, bg=WE_MED, pady=5)
        bar.pack(fill='x')
        tk.Label(bar, text=f"  {icon}  {title}",
                 font=("Segoe UI", 10, "bold"),
                 bg=WE_MED, fg=WE_WHITE).pack(side='left', padx=8)
        return bar

    def _browse_btn(self, parent, cmd):
        return tk.Button(parent, text="Browse", command=cmd,
                         font=("Segoe UI", 9, "bold"),
                         bg=WE_MED, fg=WE_WHITE,
                         activebackground=WE_DARK, activeforeground=WE_WHITE,
                         relief="flat", cursor="hand2", padx=14, pady=5)

    def _on_change(self):
        self._update_info()

    def _update_info(self):
        nt = self.network_type.get()
        sr = self.splitting_ratio.get()
        bpb = BRANCHES_PER_BOX.get(nt, 2)
        self.info_label.config(
            text=f"✔   Network: {nt}   │   Splitting Ratio: {sr}   │   "
                 f"{bpb} branches / box   │   "
                 f"{'End-Cab is calculated' if nt == 'ABF' else 'No End-Cab (Overhead)'}",
            fg=WE_MED)

    def _build_ui(self):
        hdr_outer = tk.Frame(self, bg=WE_DARK)
        hdr_outer.pack(fill='x')
        tk.Frame(hdr_outer, bg=WE_MED, height=4).pack(fill='x')

        hdr_inner = tk.Frame(hdr_outer, bg=WE_DARK)
        hdr_inner.pack(fill='x')

        badge = tk.Frame(hdr_inner, bg='#5A1E8C', padx=14, pady=8)
        badge.pack(side='left')
        tk.Label(badge, text="WE", font=("Arial", 22, "bold"), bg='#5A1E8C', fg=WE_WHITE).pack()
        tk.Frame(badge, bg='#9B6AC0', height=2, width=36).pack(fill='x')

        tk.Frame(hdr_inner, bg='#5A1E8C', width=1).pack(side='left', fill='y', pady=6)

        center = tk.Frame(hdr_inner, bg=WE_DARK)
        center.pack(side='left', fill='both', expand=True, pady=6, padx=16)
        tk.Label(center, text="Cabinet Assignment Tool — LGX", font=("Segoe UI", 17, "bold"),
                 bg=WE_DARK, fg=WE_WHITE, anchor='center').pack(fill='x')
        tk.Label(center, text="Fiber Network Planning & Optimization", font=("Segoe UI", 9),
                 bg=WE_DARK, fg='#C3A8E8', anchor='center').pack(fill='x', pady=(1, 4))
        tk.Frame(center, bg=GOLD, height=1).pack(fill='x', padx=60, pady=(0, 4))
        tk.Label(center, text="تصميم وتطوير: م/ أحمد محمد السيد", font=("Segoe UI", 9, "bold"),
                 bg=WE_DARK, fg=GOLD, anchor='center').pack(fill='x')
        tk.Label(center, text="قطاع تخطيط الشبكات غرب القاهرة", font=("Segoe UI", 8),
                 bg=WE_DARK, fg=GOLD, anchor='center').pack(fill='x', pady=(1, 0))

        tk.Frame(hdr_outer, bg='#5A1E8C', height=3).pack(fill='x')

        sub_bar = tk.Frame(self, bg='#F0EAF9', pady=3)
        sub_bar.pack(fill='x')
        tk.Label(sub_bar, text="⚡   Assigns ODF branches & LGX splitter ports to each box",
                 font=("Segoe UI", 9, "italic"), bg='#F0EAF9', fg=WE_MED).pack()

        tk.Frame(self, bg=DIVIDER, height=1).pack(fill='x', padx=20)

        body = tk.Frame(self, bg='#EDE7F6')
        body.pack(fill='both', expand=True, padx=14, pady=6)

        # ── Configuration card ──
        net_outer, net_card = _card(body, padx=0, pady=0)
        net_outer.pack(fill='x', pady=(0, 5))
        self._section_header(net_card, "⚙", "Project Configuration")

        cfg = tk.Frame(net_card, bg=CARD_BG, padx=14, pady=7)
        cfg.pack(fill='x')

        rb_style = dict(font=("Segoe UI", 10), bg=CARD_BG, fg=WE_DARK,
                        selectcolor='#D4B4F5', activebackground=CARD_BG,
                        relief='flat', bd=0, command=self._on_change)
        badge_style = dict(font=("Segoe UI", 9), bg='#F0EAF9', fg=WE_DARK,
                           padx=8, pady=2, relief='flat')

        # Network type
        tk.Label(cfg, text="Network Type:", font=("Segoe UI", 10, "bold"),
                 bg=CARD_BG, fg=WE_DARK).pack(anchor='w')
        nrow = tk.Frame(cfg, bg=CARD_BG); nrow.pack(fill='x', pady=(2, 4))
        tk.Radiobutton(nrow, text="معلق  (Overhead)", width=18, anchor="w",
                       variable=self.network_type, value='معلق', **rb_style).pack(side='left')
        tk.Label(nrow, text="2 branches / box", **badge_style).pack(side='left', padx=6)
        tk.Radiobutton(nrow, text="ABF  (حفر)", width=14, anchor="w",
                       variable=self.network_type, value='ABF', **rb_style).pack(side='left')
        tk.Label(nrow, text="4 branches / box  •  End-Cab", **badge_style).pack(side='left', padx=4)

        tk.Frame(cfg, bg=DIVIDER, height=1).pack(fill='x', pady=(4, 6))

        # Splitting ratio
        tk.Label(cfg, text="Splitting Ratio:", font=("Segoe UI", 10, "bold"),
                 bg=CARD_BG, fg=WE_DARK).pack(anchor='w')
        srow = tk.Frame(cfg, bg=CARD_BG); srow.pack(fill='x', pady=(2, 2))
        tk.Radiobutton(srow, text="1:64", width=18, anchor="w",
                       variable=self.splitting_ratio, value='1:64', **rb_style).pack(side='left')
        tk.Label(srow, text="box 1:4→SP1:16 … 1:64 direct", **badge_style).pack(side='left', padx=6)
        tk.Radiobutton(srow, text="1:32", width=14, anchor="w",
                       variable=self.splitting_ratio, value='1:32', **rb_style).pack(side='left')
        tk.Label(srow, text="1:4→SP1:8 … 1:32/1:64 direct", **badge_style).pack(side='left', padx=4)

        info_strip = tk.Frame(net_card, bg='#EDE7F6', pady=3)
        info_strip.pack(fill='x')
        self.info_label = tk.Label(info_strip, text="", font=("Segoe UI", 9, "bold"),
                                   bg='#EDE7F6', fg=WE_MED)
        self.info_label.pack()
        self._update_info()

        # ── File card ──
        file_outer, file_card = _card(body, padx=0, pady=0)
        file_outer.pack(fill='x', pady=(0, 5))
        self._section_header(file_card, "📁", "File Selection")

        file_area = tk.Frame(file_card, bg=CARD_BG, padx=14, pady=7)
        file_area.pack(fill='x')
        lbl_style = dict(font=("Segoe UI", 9, "bold"), bg=CARD_BG, fg='#555577', width=11, anchor='w')
        ent_style = dict(font=("Segoe UI", 9), bg=ENTRY_BG, relief='solid', bd=1,
                         highlightthickness=1, highlightcolor=WE_LIGHT, highlightbackground=CARD_BORDER)

        ir = tk.Frame(file_area, bg=CARD_BG); ir.pack(anchor='center', pady=3)
        tk.Label(ir, text="Input Excel:", **lbl_style).pack(side='left')
        tk.Entry(ir, textvariable=self.input_path, width=52, **ent_style).pack(side='left', padx=(6, 8), ipady=3)
        self._browse_btn(ir, self._browse_input).pack(side='left')

        or_ = tk.Frame(file_area, bg=CARD_BG); or_.pack(anchor='center', pady=3)
        tk.Label(or_, text="Output File:", **lbl_style).pack(side='left')
        tk.Entry(or_, textvariable=self.output_path, width=52, **ent_style).pack(side='left', padx=(6, 8), ipady=3)
        self._browse_btn(or_, self._browse_output).pack(side='left')

        action_row = tk.Frame(body, bg='#EDE7F6'); action_row.pack(fill='x', pady=(2, 4))
        self.run_btn = tk.Button(action_row, text="  ▶   RUN PROCESSING  ", command=self._run,
                                 font=("Segoe UI", 11, "bold"), bg=WE_DARK, fg=WE_WHITE,
                                 activebackground='#5A1E8C', activeforeground=WE_WHITE,
                                 relief="flat", cursor="hand2", padx=40, pady=6)
        self.run_btn.pack()

        style = ttk.Style(); style.theme_use('clam')
        style.configure("WE.Horizontal.TProgressbar", troughcolor='#D8CCE8',
                        background=WE_MED, darkcolor=WE_DARK, lightcolor=WE_LIGHT, borderwidth=0)
        self.progress = ttk.Progressbar(body, mode='indeterminate', style="WE.Horizontal.TProgressbar")
        self.progress.pack(fill='x', pady=(3, 0))

        log_outer = tk.Frame(body, bg=LOG_BG, bd=0); log_outer.pack(fill='both', expand=True, pady=(6, 0))
        log_hdr = tk.Frame(log_outer, bg=LOG_HDR_BG, pady=5); log_hdr.pack(fill='x')
        tk.Label(log_hdr, text="  📋  Processing Log", font=("Segoe UI", 9, "bold"),
                 bg=LOG_HDR_BG, fg='#9B6AC0').pack(side='left')
        tk.Label(log_hdr, text="console output  ●", font=("Segoe UI", 8),
                 bg=LOG_HDR_BG, fg='#00E676').pack(side='right', padx=10)

        term = tk.Frame(log_outer, bg=LOG_BG); term.pack(fill='both', expand=True)
        self.log = tk.Text(term, font=("Consolas", 9), bg=LOG_BG, fg=LOG_FG, relief="flat",
                           insertbackground=LOG_FG, selectbackground='#2A1550', padx=10, pady=8, spacing1=2)
        sb = tk.Scrollbar(term, command=self.log.yview, bg=LOG_HDR_BG, troughcolor=LOG_BG, relief='flat', width=10)
        self.log.configure(yscrollcommand=sb.set)
        self.log.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")

    def _browse_input(self):
        path = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx *.xls")])
        if path:
            self.input_path.set(path)
            base, _ = os.path.splitext(path)
            self.output_path.set(base + "_assigned.xlsx")

    def _browse_output(self):
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx")])
        if path:
            self.output_path.set(path)

    def _log(self, msg):
        self.log.insert("end", msg + "\n")
        self.log.see("end")

    def _run(self):
        inp = self.input_path.get().strip()
        out = self.output_path.get().strip()
        if not inp or not os.path.exists(inp):
            messagebox.showerror("Error", "Please select a valid input file.")
            return
        if not out:
            messagebox.showerror("Error", "Please specify an output file path.")
            return

        net_type = self.network_type.get()
        ratio    = self.splitting_ratio.get()
        self.log.delete("1.0", "end")
        self.run_btn.config(state="disabled")
        self.progress.start(10)

        def worker():
            try:
                process_file(inp, out, net_type, ratio, self._log)
                self.after(0, lambda: messagebox.showinfo("Done!", f"File saved successfully!\n{out}"))
            except Exception as e:
                self.after(0, lambda: messagebox.showerror("Error", str(e)))
                self._log(f"\n❌ Error: {e}")
            finally:
                self.after(0, self.progress.stop)
                self.after(0, lambda: self.run_btn.config(state="normal"))

        threading.Thread(target=worker, daemon=True).start()


# ── CLI entry ────────────────────────────────────────────────────────────────
def _cli():
    import sys
    def _print(m): print(m)
    args = sys.argv[1:]
    if len(args) < 2:
        print("Usage: python cab_assignment_app.py input.xlsx output.xlsx "
              "[معلق|ABF] [1:64|1:32]")
        return
    inp, out = args[0], args[1]
    net = args[2] if len(args) >= 3 else 'معلق'
    ratio = args[3] if len(args) >= 4 else '1:64'
    process_file(inp, out, net, ratio, _print)


if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1:
        _cli()
    else:
        app = App()
        app.mainloop()
